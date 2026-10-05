"""
Professional DOCX Study Guide Builder for Modern Physics classes.
Generates structured university study guides formatted with typography,
rendered mathematical equation images, tables, and callouts.
"""

from pathlib import Path
from typing import Dict, Any, List, Optional
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from ..visualization.math_renderer import MathRenderer


def set_cell_background(cell, fill_hex: str):
    """Sets background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets inner padding for a cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


class DocxStudyGuideBuilder:
    """Builds a complete, professional study guide in .docx format with rendered math."""

    NAVY = RGBColor(30, 58, 138)       # #1E3A8A
    BLUE = RGBColor(37, 99, 235)       # #2563EB
    DARK = RGBColor(30, 41, 59)        # #1E293B
    RED = RGBColor(185, 28, 28)        # #B91C1C
    AMBER = RGBColor(180, 83, 9)       # #B45309
    GREEN = RGBColor(4, 120, 87)       # #047857

    @classmethod
    def build(
        cls,
        class_name: str,
        synthesis: Dict[str, Any],
        equations: List[Dict[str, Any]],
        problems: List[Dict[str, Any]],
        segment_analyses: List[Dict[str, Any]],
        output_path: Path
    ) -> Path:
        doc = Document()
        out_file = Path(output_path).resolve()
        out_file.parent.mkdir(parents=True, exist_ok=True)
        img_cache_dir = out_file.parent / ".math_cache"
        img_cache_dir.mkdir(parents=True, exist_ok=True)

        # Page setup: Margins 1 inch
        for section in doc.sections:
            section.top_margin = Inches(1.0)
            section.bottom_margin = Inches(1.0)
            section.left_margin = Inches(1.0)
            section.right_margin = Inches(1.0)

        # Document Header
        title = synthesis.get("title") or f"Guía de Estudio — {class_name}"
        cls._add_header_block(doc, title, class_name)

        # 1. Resumen ejecutivo
        cls._add_section_heading(doc, "1. Resumen Ejecutivo")
        exec_sum = synthesis.get("executive_summary", "")
        doc.add_paragraph(exec_sum)

        # 2. Mapa conceptual (referencia y estructura)
        cls._add_section_heading(doc, "2. Mapa Conceptual y Estructura Cognitiva")
        p_map = doc.add_paragraph()
        p_map.add_run(
            "Esta clase se articula en torno a los siguientes ejes temáticos y sus relaciones directas:\n"
        )
        for dep in synthesis.get("prerequisites_and_dependencies", []):
            p = doc.add_paragraph(style="List Bullet")
            r1 = p.add_run(dep.get("concept", ""))
            r1.bold = True
            p.add_run(f" ➔ Prerrequisito para: {dep.get('needed_for', '')}")

        p_note = doc.add_paragraph()
        r_note = p_note.add_run("Nota: Puedes visualizar e interactuar con el mapa mental gráfico abriendo el archivo ")
        r_file = p_note.add_run("mapa_conceptual.html")
        r_file.bold = True
        p_note.add_run(" incluido en esta misma carpeta de resultados.")

        # 3. Conceptos fundamentales
        cls._add_section_heading(doc, "3. Conceptos Fundamentales")
        for c in synthesis.get("core_concepts", []):
            c_title = c.get("title") or c.get("concept", "")
            imp = c.get("importance", "IMPORTANTE")
            cls._add_concept_box(doc, c_title, imp, c.get("explanation") or c.get("description", ""))

        # 4. Teoría
        cls._add_section_heading(doc, "4. Teoría y Fundamentos Físicos")
        for seg in segment_analyses:
            seg_title = ", ".join(seg.get("topics", [])) or f"Segmento {seg.get('chunk_id')}"
            interval = seg.get("interval", "")
            p_sub = doc.add_heading(level=2)
            run_sub = p_sub.add_run(f"{seg_title} ")
            run_sub.font.color.rgb = cls.BLUE
            if interval:
                run_ts = p_sub.add_run(f"[Clase — {interval}]")
                run_ts.font.size = Pt(10)
                run_ts.font.color.rgb = RGBColor(100, 116, 139)

            th = seg.get("theory", "")
            if th:
                doc.add_paragraph(th)

            subtopics = seg.get("subtopics", [])
            if subtopics:
                p_st = doc.add_paragraph()
                p_st.add_run("Subtemas abordados: ").bold = True
                p_st.add_run(", ".join(subtopics))

        # 5. Ecuaciones importantes (con fórmulas renderizadas gráficamente)
        cls._add_section_heading(doc, "5. Ecuaciones Importantes")
        cls._add_equations_table(doc, equations, img_cache_dir)

        # 6. Cómo resolver problemas
        cls._add_section_heading(doc, "6. Metodología para Resolución de Problemas")
        for prob in problems:
            cls._add_problem_methodology_block(doc, prob)

        # 7. Ejemplos vistos en clase
        cls._add_section_heading(doc, "7. Ejemplos Desarrollados en Clase")
        for idx, prob in enumerate(problems, 1):
            p_ex = doc.add_heading(level=2)
            p_ex.add_run(f"Ejercicio #{idx}: {prob.get('title', '')} ")
            ts = prob.get("timestamp")
            if ts:
                r_ts = p_ex.add_run(f"[Clase — {ts}]")
                r_ts.font.size = Pt(10)
                r_ts.font.color.rgb = RGBColor(100, 116, 139)

            p_desc = doc.add_paragraph()
            p_desc.add_run("Datos suministrados: ").bold = True
            p_desc.add_run(prob.get("given_data", "") + "\n")
            p_desc.add_run("Incógnita a calcular: ").bold = True
            p_desc.add_run(prob.get("target", "") + "\n")
            p_desc.add_run("Ecuación base: ").bold = True
            p_desc.add_run(prob.get("equations_used", "") + "\n")

            doc.add_paragraph("Procedimiento detallado paso a paso:").runs[0].bold = True
            for step in prob.get("general_procedure", []):
                doc.add_paragraph(step, style="List Bullet")

        # 8. Errores comunes y advertencias
        cls._add_section_heading(doc, "8. Errores Comunes y Puntos Enfatizados por el Profesor")
        all_warnings = []
        all_errors = []
        for s in segment_analyses:
            all_warnings.extend(s.get("teacher_warnings", []))
            all_errors.extend(s.get("common_errors", []))

        if all_warnings:
            doc.add_heading("⚠️ Advertencias y Énfasis del Docente", level=2)
            for w in all_warnings:
                p_w = doc.add_paragraph(style="List Bullet")
                pts = w.get("timestamp")
                if pts:
                    p_w.add_run(f"[{pts}] ").bold = True
                p_w.add_run(w.get("point", ""))

        if all_errors:
            doc.add_heading("❌ Confusiones Recurrentes y su Aclaración", level=2)
            for err in all_errors:
                p_e = doc.add_paragraph(style="List Bullet")
                r_e = p_e.add_run("Error habitual: ")
                r_e.bold = True
                p_e.add_run(err.get("error", "") + "\n")
                r_c = p_e.add_run("Aclaración correcta: ")
                r_c.bold = True
                r_c.font.color.rgb = cls.GREEN
                p_e.add_run(err.get("correction", ""))

        # 9. Tips para examen
        cls._add_section_heading(doc, "9. Tips Concretos para Examen")
        for tip in synthesis.get("study_tips", []):
            cls._add_callout_box(doc, "💡 TIP DE EXAMEN", tip, fill_hex="FEF3C7", border_rgb=cls.AMBER)

        # 10. Preguntas potenciales de examen
        cls._add_section_heading(doc, "10. Banco de Preguntas Potenciales de Examen")
        exam_qs = synthesis.get("potential_exam_questions", {})

        doc.add_heading("A. Preguntas Conceptuales", level=2)
        for q in exam_qs.get("conceptual", []):
            doc.add_paragraph(q, style="List Bullet")

        doc.add_heading("B. Preguntas Matemáticas y Deductivas", level=2)
        for q in exam_qs.get("mathematical", []):
            doc.add_paragraph(q, style="List Bullet")

        doc.add_heading("C. Problemas de Aplicación", level=2)
        for q in exam_qs.get("problem_solving", []):
            doc.add_paragraph(q, style="List Bullet")

        # 11. Formulario (Cheatsheet)
        cls._add_section_heading(doc, "11. Formulario Compacto (Cheatsheet)")
        cls._add_cheatsheet_table(doc, equations, img_cache_dir)

        # 12. Repaso rápido
        cls._add_section_heading(doc, "12. Repaso Rápido en 5–10 Minutos")
        for pt in synthesis.get("quick_review_5min", []):
            doc.add_paragraph(pt, style="List Bullet")

        doc.save(out_file)
        return out_file

    @classmethod
    def _add_header_block(cls, doc: Document, title: str, class_name: str):
        title_p = doc.add_paragraph()
        title_run = title_p.add_run("GUÍA DE ESTUDIO UNIVERSITARIA")
        title_run.font.size = Pt(13)
        title_run.font.bold = True
        title_run.font.color.rgb = cls.BLUE
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        main_h = doc.add_heading(level=0)
        main_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        h_run = main_h.add_run(title)
        h_run.font.size = Pt(22)
        h_run.font.color.rgb = cls.NAVY

        sub_p = doc.add_paragraph()
        sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        s_run = sub_p.add_run(f"Materia: Física Moderna | Archivo: {class_name}")
        s_run.font.size = Pt(10)
        s_run.font.italic = True
        s_run.font.color.rgb = RGBColor(100, 116, 139)

        doc.add_paragraph()

    @classmethod
    def _add_section_heading(cls, doc: Document, text: str):
        h = doc.add_heading(level=1)
        run = h.add_run(text)
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = cls.NAVY

    @classmethod
    def _add_concept_box(cls, doc: Document, title: str, importance: str, text: str):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)

        if "IMPRESCINDIBLE" in importance:
            fill = "FEE2E2"
            tag_color = cls.RED
            tag = "🔴 IMPRESCINDIBLE"
        elif "IMPORTANTE" in importance:
            fill = "FEF3C7"
            tag_color = cls.AMBER
            tag = "🟡 IMPORTANTE"
        else:
            fill = "D1FAE5"
            tag_color = cls.GREEN
            tag = "🟢 COMPLEMENTARIO"

        set_cell_background(cell, fill)
        set_cell_margins(cell, top=120, bottom=120, left=160, right=160)

        p = cell.paragraphs[0]
        r_tag = p.add_run(f"{tag} — ")
        r_tag.bold = True
        r_tag.font.color.rgb = tag_color

        r_title = p.add_run(f"{title}\n")
        r_title.bold = True
        r_title.font.size = Pt(11)

        r_desc = p.add_run(text)
        r_desc.font.size = Pt(10)

        doc.add_paragraph()

    @classmethod
    def _add_callout_box(cls, doc: Document, title: str, text: str, fill_hex="F1F5F9", border_rgb=NAVY):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        set_cell_background(cell, fill_hex)
        set_cell_margins(cell, top=100, bottom=100, left=150, right=150)

        p = cell.paragraphs[0]
        r_t = p.add_run(f"{title}: ")
        r_t.bold = True
        r_t.font.color.rgb = border_rgb
        p.add_run(text)

        doc.add_paragraph()

    @classmethod
    def _add_equations_table(cls, doc: Document, equations: List[Dict[str, Any]], img_dir: Path):
        if not equations:
            doc.add_paragraph("No se registraron ecuaciones explícitas en esta clase.")
            return

        table = doc.add_table(rows=1, cols=4)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        headers = ["Ecuación / Ley", "Expresión Matemática", "Variables y Unidades", "Cuándo Usar y Condiciones"]

        hdr_cells = table.rows[0].cells
        for idx, text in enumerate(headers):
            hdr_cells[idx].text = text
            set_cell_background(hdr_cells[idx], "1E3A8A")
            p = hdr_cells[idx].paragraphs[0]
            for run in p.runs:
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.bold = True
                run.font.size = Pt(9.5)

        for eq in equations:
            row_cells = table.add_row().cells

            # Cell 1: Name & Warning & Timestamp
            p1 = row_cells[0].paragraphs[0]
            p1.add_run(eq.get("name", "")).bold = True
            ts = eq.get("timestamp")
            if ts:
                p1.add_run(f"\n[Clase — {ts}]").font.size = Pt(8.5)
            if eq.get("confidence") == "LOW" or eq.get("review_warning"):
                rw = eq.get("review_warning") or "⚠️ REVISAR"
                r_warn = p1.add_run(f"\n{rw}")
                r_warn.font.color.rgb = cls.RED
                r_warn.font.size = Pt(8.5)
                r_warn.bold = True

            # Cell 2: Rendered Math Image + LaTeX code caption
            p2 = row_cells[1].paragraphs[0]
            latex = eq.get("latex", "")
            img_path = MathRenderer.render_to_png(latex, img_dir)
            if img_path and img_path.exists():
                r_img = p2.add_run()
                r_img.add_picture(str(img_path), height=Inches(0.28))
                p2_sub = row_cells[1].add_paragraph()
                r_sub = p2_sub.add_run(f"LaTeX: {latex}")
                r_sub.font.size = Pt(7.5)
                r_sub.font.color.rgb = RGBColor(148, 163, 184)
            else:
                r2 = p2.add_run(latex)
                r2.font.name = "Consolas"
                r2.font.size = Pt(9.5)

            # Cell 3: Variables & units
            p3 = row_cells[2].paragraphs[0]
            p3.add_run(f"{eq.get('variables', '')}\n")
            r_u = p3.add_run(f"Unidades: {eq.get('units', 'SI')}")
            r_u.font.italic = True
            r_u.font.size = Pt(8.5)

            # Cell 4: When to use & conditions
            p4 = row_cells[3].paragraphs[0]
            p4.add_run(eq.get("when_to_use", ""))
            if eq.get("conditions"):
                p4.add_run(f"\nSupuestos: {eq.get('conditions')}")

            for c in row_cells:
                set_cell_margins(c, top=80, bottom=80, left=100, right=100)

        doc.add_paragraph()

    @classmethod
    def _add_problem_methodology_block(cls, doc: Document, prob: Dict[str, Any]):
        p_head = doc.add_heading(level=2)
        p_head.add_run(f"Patrón: {prob.get('title', 'Problema')} ")
        ts = prob.get("timestamp")
        if ts:
            r_ts = p_head.add_run(f"[Clase — {ts}]")
            r_ts.font.size = Pt(10)
            r_ts.font.color.rgb = RGBColor(100, 116, 139)

        tbl = doc.add_table(rows=0, cols=2)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

        fields = [
            ("¿Cómo reconocerlo?", prob.get("how_to_recognize", "")),
            ("¿Qué datos suelen aparecer?", prob.get("given_data", "")),
            ("¿Qué me están pidiendo?", prob.get("target", "")),
            ("¿Qué conceptos debo recordar?", prob.get("concepts_to_remember", "")),
            ("¿Qué ecuación/método utilizar?", prob.get("equations_used", "")),
        ]

        for label, val in fields:
            row = tbl.add_row().cells
            set_cell_background(row[0], "F1F5F9")
            p_l = row[0].paragraphs[0]
            p_l.add_run(label).bold = True
            row[1].paragraphs[0].add_run(val)
            set_cell_margins(row[0], 60, 60, 80, 80)
            set_cell_margins(row[1], 60, 60, 80, 80)

        doc.add_paragraph()
        p_proc = doc.add_paragraph()
        p_proc.add_run("Procedimiento General Paso a Paso:").bold = True
        for step in prob.get("general_procedure", []):
            doc.add_paragraph(step, style="List Bullet")

        pitfalls = prob.get("common_pitfalls", [])
        if pitfalls:
            p_err = doc.add_paragraph()
            r_err = p_err.add_run("¿Qué errores debo evitar?:")
            r_err.bold = True
            r_err.font.color.rgb = cls.RED
            for pit in pitfalls:
                doc.add_paragraph(pit, style="List Bullet")

        doc.add_paragraph()

    @classmethod
    def _add_cheatsheet_table(cls, doc: Document, equations: List[Dict[str, Any]], img_dir: Path):
        table = doc.add_table(rows=1, cols=3)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        headers = ["Nombre / Ley", "Fórmula Matemática", "Aplicación Inmediata"]
        hdr_cells = table.rows[0].cells
        for idx, text in enumerate(headers):
            hdr_cells[idx].text = text
            set_cell_background(hdr_cells[idx], "2563EB")
            p = hdr_cells[idx].paragraphs[0]
            for run in p.runs:
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.bold = True
                run.font.size = Pt(9)

        for eq in equations:
            row_cells = table.add_row().cells
            row_cells[0].paragraphs[0].add_run(eq.get("name", "")).bold = True

            p_math = row_cells[1].paragraphs[0]
            latex = eq.get("latex", "")
            img_path = MathRenderer.render_to_png(latex, img_dir)
            if img_path and img_path.exists():
                r_img = p_math.add_run()
                r_img.add_picture(str(img_path), height=Inches(0.24))
            else:
                r_eq = p_math.add_run(latex)
                r_eq.font.name = "Consolas"

            row_cells[2].paragraphs[0].add_run(eq.get("when_to_use", "")[:95] + "...")
            for c in row_cells:
                set_cell_margins(c, 60, 60, 80, 80)

        doc.add_paragraph()
